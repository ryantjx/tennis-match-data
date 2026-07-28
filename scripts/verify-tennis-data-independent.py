#!/usr/bin/env python3
"""Independent Tennis-Data/candidate comparator.

This script intentionally imports no ``open_tennis_data`` module and does not
use DuckDB. XLS/XLSX and Parquet are decoded through separate third-party
implementations from the production collector.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from collections.abc import Iterable
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any


def normalized(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def excel_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def workbook_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.casefold() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
        workbook.close()
    else:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        values = [sheet.row_values(index) for index in range(sheet.nrows)]
        workbook.release_resources()
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    return [
        {
            header: row[index] if index < len(row) else None
            for index, header in enumerate(headers)
        }
        for row in values[1:]
    ]


def source_key(tour: str, year: int, row: dict[str, Any]) -> tuple[str, ...] | None:
    played_on = excel_date(row.get("Date"))
    required = [row.get(name) for name in ("Tournament", "Winner", "Loser", "Round")]
    if (
        played_on is None
        or (
            played_on.year != year
            and not (played_on.year == year - 1 and played_on.month == 12)
        )
        or not all(normalized(value) for value in required)
    ):
        return None
    participants = sorted((normalized(row["Winner"]), normalized(row["Loser"])))
    return (
        tour,
        played_on.isoformat(),
        normalized(row["Tournament"]),
        normalized(row["Round"]),
        *participants,
    )


def source_status(row: dict[str, Any]) -> str:
    comment = normalized(row.get("Comment"))
    if comment in {"walkover", "w o", "wo"}:
        return "walkover"
    if comment in {"defaulted", "def"}:
        return "defaulted"
    if comment in {"retired", "ret"}:
        return "retired"
    if comment in {"abandoned", "abd", "abn"}:
        return "abandoned"
    return "completed"


def source_score(row: dict[str, Any]) -> str | None:
    sets: list[str] = []
    for number in range(1, 6):
        winner = row.get(f"W{number}")
        loser = row.get(f"L{number}")
        if winner in (None, "") or loser in (None, ""):
            continue
        try:
            sets.append(f"{int(float(str(winner)))}-{int(float(str(loser)))}")
        except (TypeError, ValueError):
            continue
    suffix = {
        "retired": "RET",
        "ret": "RET",
        "walkover": "W/O",
        "w o": "W/O",
        "wo": "W/O",
        "defaulted": "DEF",
        "def": "DEF",
    }.get(normalized(row.get("Comment")))
    value = " ".join((*sets, suffix) if suffix else sets).strip()
    return value or None


def source_material(tour: str, year: int, row: dict[str, Any]) -> dict[str, Any]:
    played_on = excel_date(row.get("Date"))
    return {
        "tour": tour,
        "year": year,
        "date": played_on.isoformat() if played_on else None,
        "tournament": normalized(row.get("Tournament")),
        "round": normalized(row.get("Round")),
        "participants": sorted(
            (normalized(row.get("Winner")), normalized(row.get("Loser")))
        ),
        "winner": normalized(row.get("Winner")),
        "score": source_score(row),
        "status": source_status(row),
    }


def candidate_rows(paths: Iterable[Path]) -> list[dict[str, Any]]:
    import pyarrow.parquet as parquet

    rows: list[dict[str, Any]] = []
    for path in sorted(paths):
        # Read each physical file directly. ``read_table(path)`` treats the
        # surrounding ``tour=.../year=...`` directories as Hive partitions and
        # can collide with the same typed columns stored inside the file.
        rows.extend(parquet.ParquetFile(path).read().to_pylist())
    return rows


def direct_parquet_rows(path: Path) -> list[dict[str, Any]]:
    import pyarrow.parquet as parquet

    return parquet.ParquetFile(path).read().to_pylist()


def candidate_key(row: dict[str, Any]) -> tuple[str, ...]:
    names = sorted(
        (
            normalized((row.get("player1_name") or [""])[0]),
            normalized((row.get("player2_name") or [""])[0]),
        )
    )
    played_on = row["date"]
    return (
        str(row["tour"]),
        played_on.isoformat() if hasattr(played_on, "isoformat") else str(played_on),
        normalized(row["tournament_name"]),
        normalized(row["round"]),
        *names,
    )


def candidate_material(row: dict[str, Any]) -> dict[str, Any]:
    player1_id = (row.get("player1_id") or [None])[0]
    player2_id = (row.get("player2_id") or [None])[0]
    winner_ids = set(row.get("winner_id") or [])
    winner = (
        (row.get("player1_name") or [""])[0]
        if player1_id in winner_ids
        else (row.get("player2_name") or [""])[0]
        if player2_id in winner_ids
        else ""
    )
    played_on = row.get("date")
    return {
        "tour": str(row.get("tour")),
        "year": int(row.get("year")),
        "date": played_on.isoformat() if hasattr(played_on, "isoformat") else str(played_on),
        "tournament": normalized(row.get("tournament_name")),
        "round": normalized(row.get("round")),
        "participants": sorted(
            (
                normalized((row.get("player1_name") or [""])[0]),
                normalized((row.get("player2_name") or [""])[0]),
            )
        ),
        "winner": normalized(winner),
        "score": row.get("score"),
        "status": str(row.get("status")),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-bundle", required=True, type=Path)
    parser.add_argument("--dataset", required=True, type=Path)
    parser.add_argument("--baseline-release", type=Path)
    parser.add_argument("--minimum-retention", type=float, default=0.999)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_keys: Counter[tuple[str, ...]] = Counter()
    expected_by_source_id: dict[str, dict[str, Any]] = {}
    source_ids: set[str] = set()
    source_rows = 0
    source_seasons: set[tuple[str, int]] = set()
    for path in sorted((args.source_bundle / "artifacts").glob("*.xls*")):
        match = re.fullmatch(r"(atp|wta)-(\d{4})\.xlsx?", path.name)
        if not match:
            continue
        tour, year_text = match.groups()
        year = int(year_text)
        source_seasons.add((tour, year))
        for ordinal, row in enumerate(workbook_rows(path), 2):
            source_rows += 1
            source_match_id = f"tennis-data:{tour}:{year}:{ordinal}"
            source_ids.add(source_match_id)
            key = source_key(tour, year, row)
            if key:
                expected_by_source_id[source_match_id] = source_material(tour, year, row)
    candidates = candidate_rows(args.dataset.glob("matches/tour=*/year=*/matches.parquet"))
    candidate_keys = Counter(candidate_key(row) for row in candidates)
    candidates_by_match_id = {str(row["match_id"]): row for row in candidates}
    observation_rows = candidate_rows(
        args.dataset.glob(
            "observations/tour=*/year=*/observations.parquet"
        )
    )
    source_to_match = {
        str(row["source_match_id"]): str(row["match_id"]) for row in observation_rows
    }
    for source_match_id, material in expected_by_source_id.items():
        if source_match_id in source_to_match:
            source_keys[
                (
                    str(material["tour"]),
                    str(material["date"]),
                    str(material["tournament"]),
                    str(material["round"]),
                    *material["participants"],
                )
            ] += 1
    material_differences: list[dict[str, Any]] = []
    for source_match_id, expected in expected_by_source_id.items():
        match_id = source_to_match.get(source_match_id)
        candidate = candidates_by_match_id.get(match_id or "")
        if candidate is None:
            continue
        actual = candidate_material(candidate)
        if actual != expected:
            material_differences.append(
                {
                    "source_match_id": source_match_id,
                    "expected": expected,
                    "actual": actual,
                }
            )
    missing = source_keys - candidate_keys
    extra = candidate_keys - source_keys
    candidate_seasons = {(str(row["tour"]), int(row["year"])) for row in candidates}

    disposition_records = direct_parquet_rows(
        args.dataset / "integrity/migration_dispositions.parquet"
    )
    dispositions = len(disposition_records)
    disposition_ids = [str(row["baseline_identity"]) for row in disposition_records]
    baseline_rows = 0
    retained = 0
    if args.baseline_release:
        baseline = direct_parquet_rows(args.baseline_release / "completed.parquet")
        baseline_match_ids = {str(row["match_id"]) for row in baseline}
        baseline_provenance = direct_parquet_rows(
            args.baseline_release / "provenance.parquet"
        )
        baseline_source_ids = {
            str(row["source_match_id"])
            for row in baseline_provenance
            if str(row["match_id"]) in baseline_match_ids
            and str(row.get("source_match_id", "")).startswith("tennis-data:")
        }
        baseline_rows = len(baseline_match_ids)
        retained = len(baseline_source_ids & set(source_to_match))
    retention = retained / baseline_rows if baseline_rows else 1.0
    report = {
        "source_rows": source_rows,
        "eligible_source_rows": sum(source_keys.values()),
        "candidate_rows": len(candidates),
        "disposition_rows": dispositions,
        "unmatched_source_rows": sum(missing.values()),
        "unmatched_candidate_rows": sum(extra.values()),
        "missing_seasons": sorted(source_seasons - candidate_seasons),
        "baseline_rows": baseline_rows,
        "retained_baseline_rows": retained,
        "retention": retention,
        "minimum_retention": args.minimum_retention,
        "material_field_differences": len(material_differences),
        "material_field_difference_examples": material_differences[:10],
        "duplicate_disposition_identities": dispositions - len(set(disposition_ids)),
        "unaccounted_source_identities": len(source_ids - set(disposition_ids)),
        "unknown_disposition_identities": len(set(disposition_ids) - source_ids),
        "status": "passed",
    }
    failures: list[str] = []
    if dispositions != source_rows:
        failures.append("disposition count does not equal raw source row count")
    if len(set(disposition_ids)) != dispositions:
        failures.append("a source identity has more than one disposition")
    if set(disposition_ids) != source_ids:
        failures.append("dispositions do not exactly cover source identities")
    if missing:
        failures.append("eligible Tennis-Data rows are absent from the candidate")
    if extra:
        failures.append("candidate rows lack a Tennis-Data counterpart")
    if report["missing_seasons"]:
        failures.append("one or more complete source seasons disappeared")
    if retention < args.minimum_retention:
        failures.append("baseline retention is below the required threshold")
    if material_differences:
        failures.append("candidate material fields differ from Tennis-Data")
    if failures:
        report["status"] = "failed"
        report["failures"] = failures
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, sort_keys=True))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
